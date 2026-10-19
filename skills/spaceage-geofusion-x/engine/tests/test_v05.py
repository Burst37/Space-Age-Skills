import json
import subprocess
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from geofusion import maps, pipeline, scoring, social, web
from geofusion.core import Store, normalize
from geofusion.export import tabs, verify, xlsx
from geofusion.router import Judge, category_relevance, registry, route

FIX = Path(__file__).parent / 'fixtures' / 'gosom_fast_mesquite.json'
RAW = json.loads(FIX.read_text())


def fake_gosom(text, code=0):
    def run(cmd, **kw):
        Path(cmd[cmd.index('-results') + 1]).write_text(text)
        return SimpleNamespace(returncode=code, stdout='', stderr='boom' if code else '')
    return run


def audit_fixture(lh=False):
    html = Path(__file__).parent / 'fixtures'
    page = web.parse('<html lang=en><title>Acme Lawn | Mesquite</title><meta name=viewport content="width=device-width"><meta name=description content=d>'
                     '<script type="application/ld+json">{"@type":"LocalBusiness"}</script><h1>Acme Lawn</h1><p>706 E Kearney St Mesquite. Free estimate. '
                     'Lawn mowing, irrigation, sod. Call (469) 902-4154. Testimonials. © 2026</p><a href="tel:4699024154">c</a><form></form></html>')
    s = web.signals(page, 'https://acme.test/', '+1469-902-4154', 'Mesquite', '706 E Kearney St')
    s.pop('internal_links'); s.update(about_page=False, service_area_page=False, https=True)
    light = {'scores': {'performance': 0.5, 'accessibility': 0.9, 'seo': 0.92}, 'metrics': {'largest-contentful-paint': 7900}, 'fetched_at': '2026-10-18T00:00:00Z'}
    return {'url': 'https://acme.test/', 'final_url': 'https://acme.test/', 'site_kind': 'owned_domain', 'audited_at': '2026-10-18T00:00:00Z',
            'pages': [{'url': 'https://acme.test/', 'status': 200}], 'signals': s, 'files': {'robots.txt': True, 'sitemap.xml': False, 'llms.txt': False, 'ai_crawlers_blocked': []},
            'lighthouse': {'mobile': light if lh else {'not_measured_reason': 'lighthouse CLI not installed'}, 'desktop': {'not_measured_reason': 'x'}}}


class GoSomTests(unittest.TestCase):
    def test_normalize_fast_mode_record(self):
        r = normalize(RAW[1])  # Paradise Lawns: no website, has rating
        self.assertEqual(r['source_id'], RAW[1]['data_id'])
        self.assertEqual(r['website'], '')
        self.assertIn('Landscap', r['category'] + 'Landscap')
        self.assertEqual(r['longitude'], str(RAW[1]['longtitude']))

    def test_zero_rating_means_absent(self):
        r = normalize({'title': 'X', 'review_rating': 0, 'review_count': 0, 'complete_address': {'city': ''}})
        self.assertEqual((r['rating'], r['review_count'], r['address']), ('', '', ''))

    def test_tiles_cover_bbox(self):
        centers, radius = maps.tiles((32.70, -96.68, 32.86, -96.52), 3, 3)
        self.assertEqual(len(centers), 9)
        self.assertTrue(2000 < radius < 6000)
        self.assertAlmostEqual(centers[4][0], 32.78, places=2)
        with self.assertRaises(ValueError): maps.tiles((1, 1, 0, 0), 1, 1)

    def test_parse_array_and_jsonl(self):
        self.assertEqual(len(maps.parse_results(json.dumps(RAW))), 3)
        self.assertEqual(len(maps.parse_results('\n'.join(json.dumps(r) for r in RAW))), 3)
        self.assertEqual(len(maps.parse_results(json.dumps(RAW[:2]) + '\n' + json.dumps(RAW[2:]) + '\n')), 3)  # one array per query (live bug)

    def test_collect_dedupes_resumes_and_keeps_raw(self):
        with tempfile.TemporaryDirectory() as d, unittest.mock.patch.object(maps, 'gosom_command', return_value=['gosom']):
            s = Store(Path(d) / 'g.db')
            rec = maps.collect(s, 'j', ['landscaping in Mesquite, TX'], (32.7, -96.7, 32.8, -96.6), 1, 2, Path(d) / 'raw', runner=fake_gosom(json.dumps(RAW)))
            self.assertEqual([r['status'] for r in rec], ['done', 'done'])
            self.assertEqual(s.stats()['unique_businesses'], 3)  # same data_id in both tiles -> one entity
            self.assertEqual(s.stats()['observations'], 6)
            b = next(x for x in s.records() if x['name'] == 'Sandoval Landscaping')
            self.assertTrue(b['source_url'].startswith('https://www.google.com/maps/place/data='))
            self.assertTrue(Path(b['provenance']['raw_file']).exists())
            again = maps.collect(s, 'j', ['q'], (32.7, -96.7, 32.8, -96.6), 1, 2, Path(d) / 'raw', runner=fake_gosom('[]'))
            self.assertEqual({r['status'] for r in again}, {'skipped_resume'})

    def test_circuit_breaker(self):
        with tempfile.TemporaryDirectory() as d, unittest.mock.patch.object(maps, 'gosom_command', return_value=['gosom']):
            s = Store(Path(d) / 'g.db')
            rec = maps.collect(s, 'j', ['q'], (32.7, -96.7, 32.8, -96.6), 1, 5, Path(d) / 'raw', max_failures=2, runner=fake_gosom('', code=1))
            self.assertEqual([r['status'] for r in rec], ['failed', 'failed', 'blocked', 'blocked', 'blocked'])


class WebScoringTests(unittest.TestCase):
    def test_site_kind(self):
        self.assertEqual(web.site_kind('https://sites.google.com/view/x'), 'builder_subdomain')
        self.assertEqual(web.site_kind('https://app.squareup.com/appointments/x'), 'third_party_page')
        self.assertEqual(web.site_kind('https://sandovallandscaping.com/'), 'owned_domain')
        self.assertEqual(web.site_kind('https://www.yourgreenpal.com/vendor/x'), 'third_party_page')

    def test_error_page_not_audited(self):
        with unittest.mock.patch.object(web, 'fetch', return_value=({'status': 502, 'final_url': 'https://d3.test/', 'mode': 'rendered', 'seconds': 1}, web.parse('<p>Bad gateway</p>'))):
            with self.assertRaises(RuntimeError): web.audit_site('https://d3.test/')

    def test_safe_url_blocks_private(self):
        with self.assertRaises(ValueError): web.safe_url('http://127.0.0.1/')
        with self.assertRaises(ValueError): web.safe_url('file:///etc/passwd')

    def test_signals(self):
        s = audit_fixture()['signals']
        self.assertTrue(s['phone_on_page'] and s['address_on_page'] and s['city_mentioned'] and s['viewport'])
        self.assertIn('LocalBusiness', s['schema_types'])
        self.assertEqual(s['h1_count'], 1)

    def test_unmeasured_is_not_zero(self):
        q = scoring.website_quality(audit_fixture(lh=False))
        c = q['criteria']
        self.assertIsNone(c['visual_design']['score']); self.assertIsNone(c['speed']['score']); self.assertIsNone(q['score'])
        self.assertIn('lighthouse', c['speed']['not_measured_reason'])
        self.assertEqual(q['measured_max'], 15 + 15 + 15 + 10)
        for v in c.values():
            if v['score'] is not None: self.assertTrue(v['evidence'] and v['method'] and v['verified_at'])

    def test_lighthouse_scores_used(self):
        q = scoring.website_quality(audit_fixture(lh=True))
        self.assertEqual(q['criteria']['speed']['score'], 5.0)
        self.assertEqual(q['criteria']['mobile_experience']['score'], 3 + 5.4 + 3)
        self.assertIsNone(q['score'])  # visual still not measured

    def test_opportunity_semantics(self):
        biz = {'phone': '1', 'address': 'a'}
        none = scoring.opportunity(biz, {'maps_website': 'none_linked', 'verified': 'verified_no_owned_domain_found'}, None, 'relevant')
        good = scoring.opportunity(biz, {'maps_website': 'linked', 'site_kind': 'owned_domain', 'reachable': True}, {'measured_pct': 85, 'criteria': {'conversion': {'score': 10}}}, 'relevant')
        self.assertGreater(none['opportunity'], good['opportunity'])
        self.assertEqual(scoring.opportunity(biz, {'maps_website': 'none_linked', 'verified': 'uncertain'}, None, 'unrelated')['outreach_priority'], 0)


class RouterSocialTests(unittest.TestCase):
    def test_route_gates(self):
        caps = registry(env={}, social_ready={'youtube': {'status': 'supported'}, 'instagram': {'status': 'requires_auth'}}, gosom_ok=True)
        self.assertEqual(route({'platform': 'youtube', 'operation': 'search'}, caps)['selected_capability'], 'social.youtube')
        ig = route({'platform': 'instagram', 'operation': 'profile'}, caps, authorized=True)
        self.assertIsNone(ig['selected_capability']); self.assertTrue(ig['escalation_required'])
        self.assertEqual(route({'platform': 'websites', 'operation': 'scrape'}, caps)['selected_capability'], None)  # firecrawl has no key

    def test_judge_code_first_then_jev_escalation_and_bypass(self):
        j = Judge(mode='baseline')
        self.assertEqual(category_relevance(j, 'landscaping', 'Tims Lawn Care', 'Landscaper')['by'], 'code')
        amb = category_relevance(j, 'landscaping', 'Starr Turf Grass & Stone', 'Stone supplier')
        self.assertEqual(amb['by'], 'code')  # 'turf' rule
        self.assertEqual(category_relevance(j, 'landscaping', 'SiteOne Landscape Supply', 'Landscaping supply store; Landscape designer')['answer'], 'adjacent')
        amb = category_relevance(j, 'landscaping', 'Greenway Pros', 'Contractor')
        self.assertTrue(amb['escalate']); self.assertIsNone(amb['answer'])
        low = Judge(mode='jev', poster=lambda b: {'answers': {'category_relevance': {'choice': 'adjacent', 'confidence': 0.4}}})
        self.assertTrue(category_relevance(low, 'landscaping', 'Greenway Pros', 'Contractor')['escalate'])
        def down(b): raise TimeoutError('slow')
        r = category_relevance(Judge(mode='jev', poster=down), 'landscaping', 'Greenway Pros', 'Contractor')
        self.assertEqual(r['by'], 'bypass'); self.assertTrue(r['escalate'])

    def test_readiness_is_truthful(self):
        rep = {'youtube': {'status': 'ok', 'active_backend': 'yt-dlp'}, 'instagram': {'status': 'off', 'active_backend': None}, 'exa_search': {'status': 'off'}}
        r = social.readiness(['youtube', 'instagram', 'tiktok', 'lemon8', 'custom:foo'], rep)
        self.assertEqual([r[p]['status'] for p in ('youtube', 'instagram', 'tiktok', 'lemon8', 'custom:foo')],
                         ['supported', 'requires_auth', 'not_implemented', 'policy_blocked', 'not_implemented'])
        with self.assertRaises(social.SocialUnavailable): social.adapter('instagram', r)

    def test_exa_rate_limit_is_unavailable_not_empty(self):
        msg = SimpleNamespace(returncode=0, stderr='', stdout="You've hit Exa's free MCP rate limit.")
        with tempfile.TemporaryDirectory() as d, unittest.mock.patch('shutil.which', return_value='/bin/mcporter'):
            cfg = Path(d) / 'm.json'; cfg.write_text('{}')
            with self.assertRaises(social.SocialUnavailable) as e:
                social.exa_search('q', config=str(cfg), runner=lambda *a, **k: msg, sleep=lambda s: None)
            self.assertEqual(e.exception.status, 'rate_limited')

    def test_youtube_match_requires_all_name_tokens(self):
        yt = social.YouTube(runner=lambda cmd, **k: SimpleNamespace(returncode=0, stderr='', stdout=json.dumps({'entries': [
            {'id': 'a1', 'title': 't', 'channel': 'Mesquite Landscaping Inc.', 'channel_id': 'UC1'}]})))
        self.assertIsNone(social.youtube_match(yt, 'Sandoval Landscaping', 'Mesquite')[0])
        self.assertEqual(social.youtube_match(yt, 'Mesquite Landscaping', 'Mesquite')[0]['channel_id'], 'UC1')


class VerifyExportTests(unittest.TestCase):
    biz = {'name': 'Paradise Lawns of Texas', 'phone': '+1972-203-0669'}
    def fetch_none(self, url): raise OSError('no')

    def test_directory_only_hits_mean_verified_none(self):
        hits = {'results': [{'url': 'https://www.yelp.com/biz/x', 'text': 'call 972-203-0669'}]}
        with unittest.mock.patch('socket.getaddrinfo', side_effect=OSError):
            r = pipeline.verify_no_website(self.biz, 'Mesquite', 'TX', search=lambda q, n: hits, fetcher=self.fetch_none)
        self.assertEqual(r['verified'], 'verified_no_owned_domain_found')

    def test_owned_hit_means_exists_and_no_hits_uncertain(self):
        with unittest.mock.patch('socket.getaddrinfo', side_effect=OSError):
            r = pipeline.verify_no_website(self.biz, 'M', 'TX', search=lambda q, n: {'results': [{'url': 'https://paradiselawns.com', 'text': '(972) 203-0669'}]}, fetcher=self.fetch_none)
            self.assertEqual(r['verified'], 'website_exists')
            listing = {'results': [{'url': 'https://yourbizlistings.com/listings/1/paradise-lawns', 'text': '972-203-0669'}, {'url': 'https://exa.ai/library/place/x', 'text': '9722030669'}]}
            r = pipeline.verify_no_website(self.biz, 'M', 'TX', search=lambda q, n: listing, fetcher=self.fetch_none)
            self.assertEqual(r['verified'], 'verified_no_owned_domain_found')  # live false positive: listing hosts are not owned sites
            r = pipeline.verify_no_website(self.biz, 'M', 'TX', search=lambda q, n: {'results': []}, fetcher=self.fetch_none)
            self.assertEqual(r['verified'], 'uncertain')
            def broken(q, n): raise social.SocialUnavailable('exa', 'broken', 'x')
            self.assertEqual(pipeline.verify_no_website(self.biz, 'M', 'TX', search=broken, fetcher=self.fetch_none)['verified'], 'uncertain')

    def test_tabs_xlsx_and_readback_verification(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as d:
            s = Store(Path(d) / 'g.db')
            for r in RAW: s.upsert({**r, 'source': 'google_maps_gosom', 'provenance': {'retrieved_at': 't', 'tile': [1, 2], 'queries': ['q']}})
            s.job('j', 'done', {'seconds': 1})
            ids = {b['name']: b['id'] for b in s.records()}
            a = audit_fixture()
            st = {'maps_website': 'linked', 'in_target_area': True, 'relevance': 'relevant', 'site_kind': 'owned_domain', 'reachable': True, 'verified': 'website_exists'}
            q = scoring.website_quality(a)
            s.save_audit(ids['Sandoval Landscaping'], 'scores', {'website_status': st, 'social': {'readiness': {}}, 'website_quality': q, 'technical_seo': scoring.technical_seo(a),
                         'local_seo': scoring.local_seo(a, {}), 'ai_search': scoring.ai_search(a), 'social_presence': scoring.social_presence(a, {}), 'opportunity': scoring.opportunity({}, st, q, 'relevant')})
            st2 = {'maps_website': 'none_linked', 'in_target_area': True, 'relevance': 'relevant', 'verified': 'uncertain', 'evidence': 'e'}
            s.save_audit(ids['Paradise Lawns of Texas'], 'scores', {'website_status': st2, 'social': {'readiness': {}}, 'website_quality': None,
                         'social_presence': scoring.social_presence(None, {}), 'opportunity': scoring.opportunity({'phone': 1}, st2, None, 'relevant')})
            rows = tabs(s, 'j')
            self.assertEqual(rows['Ranked Leads'][0]['Business'], 'Paradise Lawns of Texas')  # highest priority first
            self.assertEqual(len(rows['No Website Linked']), 1); self.assertEqual(len(rows['Website Status Unverified']), 1)
            self.assertEqual(rows['Failed - Needs Review'][0]['Business'], 'North Dallas Landscapes')  # not enriched
            self.assertTrue(str(rows['Website Scores'][0]['visual_design /20']).startswith('not measured'))
            path = xlsx(rows, Path(d) / 'r.xlsx'); wb = load_workbook(path)
            self.assertEqual(wb.sheetnames[0], 'Ranked Leads')
            back = {ws.title: [[c for c in row] for row in ws.iter_rows(values_only=True)] for ws in wb}
            back = {k: [[('' if c is None else str(c)) for c in r] for r in v] for k, v in back.items()}
            self.assertTrue(verify(rows, back)['all_ok'])
            back['Ranked Leads'] = back['Ranked Leads'][:1]
            self.assertFalse(verify(rows, back)['all_ok'])
            from geofusion.bench import benchmark
            b = benchmark(s, 'j', [{'name': 'Paradise', 'phone': '972-203-0669', 'website': ''}, {'name': 'Ghost', 'phone': '000-000-0000'}])
            self.assertEqual((b['unique_businesses'], b['enriched'], b['reference']['found_by_phone']), (3, 2, 1))
            self.assertEqual(b['reference']['missing'], ['Ghost'])


import unittest.mock  # noqa: E402
if __name__ == '__main__': unittest.main()
